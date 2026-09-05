"""Document text extraction with layered fallbacks.

Order of attempts, stopping at the first that yields usable text:
    1. Tika server  — handles PDF, DOCX, DOC, RTF, ODT, HTML, TXT.
    2. Pure-Python  — pypdf / python-docx, so the pipeline still runs when no
                      JVM is available (developer laptops, CI).
    3. OCR          — for scans and images that yield almost no characters.

A document that survives all three with no text is never dropped: it comes back
with `char_count == 0` and warnings attached, and the pipeline routes it to the
manual-review bucket.
"""

from __future__ import annotations

import hashlib
import logging
import mimetypes
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from rescan.config import settings
from rescan.extract.ocr import OcrUnavailable, ocr_file
from rescan.extract.tika import TikaClient, TikaUnavailable
from rescan.schemas import ExtractionResult

log = logging.getLogger(__name__)

TEXT_SUFFIXES = {".txt", ".md", ".text", ".csv"}


def content_hash(data: bytes) -> str:
    """Stable hash used to dedup re-runs of the same batch during development."""
    return hashlib.sha256(data).hexdigest()


def _guess_content_type(path: Path) -> str | None:
    guessed, _ = mimetypes.guess_type(path.name)
    return guessed


def _extract_pdf_pypdf(data: bytes) -> tuple[str, int | None]:
    try:
        import io

        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        pages = [(page.extract_text() or "") for page in reader.pages]
        return "\n\n".join(pages), len(reader.pages)
    except Exception as exc:  # pypdf raises a wide range on malformed files
        log.debug("pypdf failed: %s", exc)
        return "", None


def _extract_docx(data: bytes) -> str:
    try:
        import io

        import docx

        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append("\t".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    except Exception as exc:
        log.debug("python-docx failed: %s", exc)
        return ""


class Extractor:
    """Thread-pooled extraction over a batch of files."""

    def __init__(self, tika: TikaClient | None = None) -> None:
        self._tika = tika if tika is not None else TikaClient()
        self._tika_ok: bool | None = None

    def _tika_usable(self) -> bool:
        if self._tika_ok is None:
            self._tika_ok = self._tika.is_available()
            if not self._tika_ok:
                log.warning(
                    "Tika server unavailable at %s; falling back to pure-Python extraction",
                    self._tika.base_url,
                )
        return self._tika_ok

    def extract_bytes(self, data: bytes, filename: str) -> ExtractionResult:
        path = Path(filename)
        suffix = path.suffix.lower()
        content_type = _guess_content_type(path)
        result = ExtractionResult(content_type=content_type)

        text = ""

        # --- 1. Tika ---
        if self._tika_usable():
            try:
                text = self._tika.extract_text(data, content_type)
                result.backend = "tika"
            except TikaUnavailable as exc:
                result.warnings.append(f"tika: {exc}")
                # A per-document failure does not mean the server is gone, so
                # only a connection-level failure clears the cached health flag.
                if "unreachable" in str(exc):
                    self._tika_ok = None

        # --- 2. Pure-Python ---
        if not text.strip():
            if suffix == ".pdf":
                text, pages = _extract_pdf_pypdf(data)
                result.page_count = pages
                if text.strip():
                    result.backend = "pypdf"
            elif suffix == ".docx":
                text = _extract_docx(data)
                if text.strip():
                    result.backend = "docx"
            elif suffix in TEXT_SUFFIXES or content_type == "text/plain":
                text = data.decode("utf-8", errors="replace")
                if text.strip():
                    result.backend = "plaintext"

        # --- 3. OCR ---
        if settings.ocr_enabled and len(text.strip()) < settings.ocr_char_threshold:
            ocr_text = self._try_ocr(data, path, result)
            if len(ocr_text.strip()) > len(text.strip()):
                text = ocr_text
                result.backend = "ocr"
                result.ocr_used = True

        result.text = text.strip()
        result.char_count = len(result.text)
        if result.char_count == 0:
            result.warnings.append("no text recovered by any backend")
        return result

    def _try_ocr(self, data: bytes, path: Path, result: ExtractionResult) -> str:
        import tempfile

        try:
            with tempfile.NamedTemporaryFile(suffix=path.suffix or ".pdf", delete=True) as tmp:
                tmp.write(data)
                tmp.flush()
                return ocr_file(Path(tmp.name))
        except OcrUnavailable as exc:
            result.warnings.append(f"ocr unavailable: {exc}")
        except Exception as exc:
            result.warnings.append(f"ocr failed: {exc}")
        return ""

    def extract_path(self, path: Path) -> ExtractionResult:
        return self.extract_bytes(path.read_bytes(), path.name)

    def extract_many(self, paths: list[Path], workers: int | None = None) -> list[ExtractionResult]:
        workers = workers or settings.extract_workers
        with ThreadPoolExecutor(max_workers=workers) as pool:
            return list(pool.map(self.extract_path, paths))

    def close(self) -> None:
        self._tika.close()
